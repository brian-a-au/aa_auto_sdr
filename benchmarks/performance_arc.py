"""Offline performance arc comparisons, using one interpreter/dependency set.

Example: uv run --no-sync python benchmarks/performance_arc.py --baseline-root /tmp/base
Warm operations exclude imports, fixture construction and authentication. VRS
stats/inventory include command credential/resolve orchestration with stub auth.
Resolution is a local preloaded-list operation, exports cover the real single
pipeline with a stub SDK, and trending reads real local snapshots.
Cold CLI startup is reported separately. No HTTP or live credentials are used.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import os
import platform
import statistics
import subprocess
import sys
import tempfile
import time
import tracemalloc
from dataclasses import asdict, replace
from datetime import UTC, datetime, timedelta
from importlib.metadata import version
from pathlib import Path

from vrs_invocation_sharing import WORKLOADS, measure, source_identity

CASES = (*WORKLOADS, "resolution-exact", "resolution-name", "resolution-fuzzy", "excel", "json-snapshot", "trending")


def peak_rss():
    try:
        import resource
    except ImportError:
        return None
    return resource.getrusage(resource.RUSAGE_SELF).ru_maxrss * (1 if sys.platform == "darwin" else 1024)


def digest(value):
    return hashlib.sha256(json.dumps(value, sort_keys=True, default=str).encode()).hexdigest()


def local_measure(args):
    sys.path.insert(0, str(args.source_root / "src"))
    import pandas as pd

    from aa_auto_sdr.api.client import AaClient
    from aa_auto_sdr.api.fetch import resolve_rsid
    from aa_auto_sdr.api.models import ReportSuite, Segment
    from aa_auto_sdr.pipeline.single import run_single
    from aa_auto_sdr.sdr.document import SdrDocument
    from aa_auto_sdr.snapshot.store import save_snapshot
    from aa_auto_sdr.snapshot.trending import WindowSpec, compute_trending

    now = datetime(2026, 1, 1, tzinfo=UTC)
    suites = [
        {"rsid": f"bench.{i}", "name": hashlib.sha256(str(i).encode()).hexdigest()[:20]} for i in range(args.suites)
    ]
    doc = SdrDocument(
        report_suite=ReportSuite("bench.0", "Benchmark", "UTC", "USD", None),
        dimensions=[],
        metrics=[],
        calculated_metrics=[],
        virtual_report_suites=[],
        classifications=[],
        segments=[
            Segment(
                f"s{i}",
                f"Segment {i}",
                "Nested definition",
                "bench.0",
                1,
                {"container": {"preds": [{"func": "eq", "value": j} for j in range(args.depth)]}},
            )
            for i in range(args.components)
        ],
        captured_at=now,
        tool_version="benchmark",
    )
    fixture = json.loads((args.source_root / "tests/fixtures/sample_rs.json").read_text())
    fixture["segments"] = [
        {**fixture["segments"][0], "id": segment.id, "definition": segment.definition} for segment in doc.segments
    ]
    records = dict(
        zip(
            (
                "getReportSuites",
                "getDimensions",
                "getMetrics",
                "getSegments",
                "getCalculatedMetrics",
                "getClassificationDatasets",
                "getVirtualReportSuites",
            ),
            (
                [fixture["report_suite"]],
                fixture["dimensions"],
                fixture["metrics"],
                fixture["segments"],
                fixture["calculated_metrics"],
                fixture["classification_datasets"],
                [],
            ),
            strict=True,
        )
    )

    class Handle:
        vrs_calls = 0
        listing_calls = 0

        def __getattr__(self, name):
            def fetch(*_args, **_kwargs):
                if name == "getReportSuites":
                    self.listing_calls += 1
                if name == "getVirtualReportSuites":
                    self.vrs_calls += 1
                return pd.DataFrame(records[name])

            return fetch

    client = AaClient(handle=Handle(), company_id="benchmark")
    original = digest(doc.to_dict())
    with tempfile.TemporaryDirectory(prefix="aa-arc-") as folder:
        root = Path(folder)
        if args.workload == "trending":
            for i in range(args.snapshots):
                save_snapshot(replace(doc, captured_at=now + timedelta(hours=i)), snapshot_dir=root)

        def operation():
            if args.workload.startswith("resolution-"):
                mode = args.workload.removeprefix("resolution-")
                identifiers = [r["rsid"] if mode == "exact" else r["name"] for r in suites]
                if mode == "fuzzy":
                    identifiers = [name + "x" for name in identifiers]
                return [
                    resolve_rsid(
                        None, name, name_match="fuzzy" if mode == "fuzzy" else "insensitive", preloaded_suites=suites
                    )
                    for name in identifiers
                ]
            if args.workload == "trending":
                return asdict(
                    compute_trending(
                        snapshot_dir=root,
                        rsid="bench.0",
                        window=WindowSpec("30d", now - timedelta(days=1), now + timedelta(days=30)),
                    )
                )
            return run_single(
                client=client,
                rsid="demo.prod",
                formats=["excel" if args.workload == "excel" else "json"],
                output_dir=root / "output",
                captured_at=now,
                tool_version="benchmark",
                snapshot_dir=root / "snapshots" if args.workload == "json-snapshot" else None,
            )

        operation()
        if args.allocation:
            tracemalloc.start()
        client.handle.vrs_calls = 0
        client.handle.listing_calls = 0
        started = time.perf_counter()
        value = operation()
        seconds = time.perf_counter() - started
        rss = peak_rss()
        peak = tracemalloc.get_traced_memory()[1] if args.allocation else None
        if args.allocation:
            tracemalloc.stop()
        if args.workload == "excel":
            from openpyxl import load_workbook

            book = load_workbook(root / "output/demo.prod.xlsx")
            value = {
                sheet.title: {
                    "cells": [
                        [
                            (
                                c.value,
                                c.data_type,
                                c.number_format,
                                c.style_id,
                                c.hyperlink.target if c.hyperlink else None,
                            )
                            for c in row
                        ]
                        for row in sheet
                    ],
                    "widths": {name: col.width for name, col in sheet.column_dimensions.items()},
                    "row_heights": {str(number): row.height for number, row in sheet.row_dimensions.items()},
                    "filter": sheet.auto_filter.ref,
                    "freeze_panes": sheet.freeze_panes,
                }
                for sheet in book
            }
            book.close()
        elif args.workload == "json-snapshot":
            value = {str(p.relative_to(root)): json.loads(p.read_text()) for p in sorted(root.rglob("*.json"))}
        assert digest(doc.to_dict()) == original, "Fixture mutated"
        return {
            "seconds": seconds,
            "peak_bytes": peak,
            "sdk_vrs_calls": client.handle.vrs_calls,
            "sdk_listing_calls": client.handle.listing_calls,
            "process_peak_rss_bytes": rss,
            "output_sha256": digest(value),
        }


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--baseline-root", type=Path)
    parser.add_argument("--candidate-root", type=Path, default=Path.cwd())
    parser.add_argument("--cases", nargs="+", choices=CASES, default=list(CASES))
    parser.add_argument("--rounds", type=int, default=7)
    parser.add_argument(
        "--suites", type=int, default=50, help="Use 1, 50 or 500 for small/organization/stress fixtures"
    )
    parser.add_argument("--vrs-per-suite", type=int, default=20)
    parser.add_argument("--components", type=int, default=1000)
    parser.add_argument("--depth", type=int, default=20, help="Predicates per nested segment definition")
    parser.add_argument("--snapshots", type=int, default=10)
    parser.add_argument("--workers", type=int, default=8)
    parser.add_argument("--output", type=Path)
    parser.add_argument("--source-root", type=Path, help=argparse.SUPPRESS)
    parser.add_argument("--workload", choices=CASES, help=argparse.SUPPRESS)
    parser.add_argument("--allocation", action="store_true", help=argparse.SUPPRESS)
    args = parser.parse_args()
    args.latency_ms = 0
    args.rss_reader = peak_rss
    if args.source_root:
        result = measure(args) if args.workload in WORKLOADS else local_measure(args)
        result.pop("simulated_http_pages", None)
        print(json.dumps(result))
        return
    if not args.baseline_root or args.rounds < 7:
        parser.error("--baseline-root and at least seven paired rounds are required")
    roots = {"baseline": args.baseline_root.resolve(), "candidate": args.candidate_root.resolve()}
    identities = {key: source_identity(root) for key, root in roots.items()}
    if len({v["lock_sha256"] for v in identities.values()}) != 1:
        parser.error("baseline and candidate must use identical lockfiles")
    for key, root in roots.items():
        identities[key]["git_sha"] = (
            subprocess.check_output(["git", "-C", str(root), "rev-parse", "HEAD"], text=True).strip()
            if (root / ".git").exists()
            else (root / ".benchmark-source-sha").read_text().strip()
        )
    report = {
        "python": sys.version,
        "platform": platform.platform(),
        "dependencies": {name: version(name) for name in ("aanalytics2", "pandas", "numpy", "openpyxl", "xlsxwriter")},
        "identities": identities,
        "fixtures": {
            k: getattr(args, k) for k in ("suites", "vrs_per_suite", "components", "depth", "snapshots", "workers")
        },
        "scope": __doc__,
        "rounds": args.rounds,
        "cases": {},
        "cold_cli_seconds": {},
    }
    for key, root in roots.items():
        start = time.perf_counter()
        subprocess.run(
            [sys.executable, "-m", "aa_auto_sdr", "--version"],
            env={**os.environ, "PYTHONPATH": str(root / "src")},
            check=True,
            capture_output=True,
        )
        report["cold_cli_seconds"][key] = time.perf_counter() - start
    for case in args.cases:
        samples = {key: [] for key in roots}
        expected = None
        for allocation in (False, True):
            for index in range(args.rounds if not allocation else 1):
                for key in ("baseline", "candidate") if index % 2 == 0 else ("candidate", "baseline"):
                    command = [
                        sys.executable,
                        str(Path(__file__).resolve()),
                        "--source-root",
                        str(roots[key]),
                        "--workload",
                        case,
                    ]
                    for field in ("suites", "vrs_per_suite", "components", "depth", "snapshots", "workers"):
                        command.extend(["--" + field.replace("_", "-"), str(getattr(args, field))])
                    if allocation:
                        command.append("--allocation")
                    value = json.loads(subprocess.check_output(command, text=True))
                    signature = (value["output_sha256"], value["sdk_vrs_calls"], value["sdk_listing_calls"])
                    expected = expected or signature
                    assert signature == expected, (case, key, "semantic or SDK enumeration mismatch")
                    value["allocation_run"] = allocation
                    samples[key].append(value)
        summary = {}
        for key, values in samples.items():
            wall = [v["seconds"] for v in values if not v["allocation_run"]]
            summary[key] = {
                "median_seconds": statistics.median(wall),
                "min_seconds": min(wall),
                "max_seconds": max(wall),
                "peak_python_bytes": values[-1]["peak_bytes"],
                "process_peak_rss_bytes": values[-1]["process_peak_rss_bytes"],
                "sdk_vrs_enumerations": values[0]["sdk_vrs_calls"],
                "sdk_report_suite_enumerations": values[0]["sdk_listing_calls"],
            }
        report["cases"][case] = {"summary": summary, "samples": samples}
        print(f"{case}: {summary}", file=sys.stderr, flush=True)
    rendered = json.dumps(report, indent=2) + "\n"
    if args.output:
        args.output.write_text(rendered)
    else:
        print(rendered)


if __name__ == "__main__":
    main()
