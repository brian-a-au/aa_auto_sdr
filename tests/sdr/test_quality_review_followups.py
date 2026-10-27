"""U10 review follow-ups: auto-enable for naming rules, unmatched exceptions,
CSV report neutralization, Title Case small words, capped ID lists, the
case_inconsistency message, and the Excel summary layout."""

from __future__ import annotations

import argparse
import csv as _csv
import logging
from dataclasses import replace
from datetime import UTC, datetime
from pathlib import Path
from typing import Any

import pytest
from openpyxl import load_workbook

from aa_auto_sdr.api import models
from aa_auto_sdr.cli import main as cli_main
from aa_auto_sdr.output.coverage import QUALITY_ISSUE_COLUMNS, quality_issue_cells
from aa_auto_sdr.output.writers.excel import ExcelWriter
from aa_auto_sdr.sdr.document import COMPONENT_FAMILIES, SdrDocument
from aa_auto_sdr.sdr.quality import (
    Issue,
    NamingRule,
    QualityException,
    QualityRules,
    SeverityLevel,
    matches_case_style,
    run_audits,
)
from aa_auto_sdr.sdr.quality_policy import write_quality_report

_CLI_MAIN_LOGGER = "aa_auto_sdr.cli.main"
_QUALITY_LOGGER = "aa_auto_sdr.sdr.quality"


def _dim(cid: str, name: str) -> models.Dimension:
    return models.Dimension(
        id=cid, name=name, type="string", category=None, parent="", pathable=False, description=None
    )


def _doc(dims: list[Any]) -> SdrDocument:
    return SdrDocument(
        report_suite=models.ReportSuite(rsid="rs1", name="RS1", timezone=None, currency=None, parent_rsid=None),
        dimensions=dims,
        metrics=[],
        segments=[],
        calculated_metrics=[],
        virtual_report_suites=[],
        classifications=[],
        captured_at=datetime(2026, 1, 1, tzinfo=UTC),
        tool_version="test",
        collected_components=COMPONENT_FAMILIES,
    )


_NAMING = QualityRules(naming={"dimensions": NamingRule(case="snake_case")})


class TestAutoEnableForNamingRules:
    def _ns(self, **kw: Any) -> argparse.Namespace:
        base = {"quality_report": None, "fail_on_quality": None, "audit_naming": False, "flag_stale": False}
        base.update(kw)
        return argparse.Namespace(**base)

    def test_policy_with_only_naming_enables_naming_audit(self, caplog: pytest.LogCaptureFixture) -> None:
        ns = self._ns(quality_rules=_NAMING)
        with caplog.at_level(logging.INFO, logger=_CLI_MAIN_LOGGER):
            cli_main._apply_quality_auto_enable(ns)
        assert ns.audit_naming is True
        assert ns.flag_stale is False
        (event,) = [r for r in caplog.records if r.getMessage().startswith("quality_auto_enabled")]
        assert getattr(event, "reason", None) == "policy_naming"
        assert getattr(event, "audit_naming", None) is True

    def test_flag_stale_with_gate_keeps_naming_rules(self, caplog: pytest.LogCaptureFixture) -> None:
        ns = self._ns(flag_stale=True, fail_on_quality="HIGH", quality_rules=_NAMING)
        with caplog.at_level(logging.INFO, logger=_CLI_MAIN_LOGGER):
            cli_main._apply_quality_auto_enable(ns)
        assert ns.audit_naming is True
        assert ns.flag_stale is True
        (event,) = [r for r in caplog.records if r.getMessage().startswith("quality_auto_enabled")]
        assert getattr(event, "reason", None) == "policy_naming"

    def test_quality_flags_path_carries_reason(self, caplog: pytest.LogCaptureFixture) -> None:
        ns = self._ns(quality_report="json")
        with caplog.at_level(logging.INFO, logger=_CLI_MAIN_LOGGER):
            cli_main._apply_quality_auto_enable(ns)
        (event,) = [r for r in caplog.records if r.getMessage().startswith("quality_auto_enabled")]
        assert getattr(event, "reason", None) == "quality_flags"

    def test_exceptions_only_do_not_enable(self, caplog: pytest.LogCaptureFixture) -> None:
        rules = QualityRules(
            exceptions=(QualityException(rule_id="dq.stale_keyword.v1", component_id="x", reason="r"),)
        )
        ns = self._ns(quality_rules=rules)
        cli_main._apply_quality_auto_enable(ns)
        assert ns.audit_naming is False

    def test_already_on_does_not_log(self, caplog: pytest.LogCaptureFixture) -> None:
        ns = self._ns(audit_naming=True, quality_rules=_NAMING)
        with caplog.at_level(logging.INFO, logger=_CLI_MAIN_LOGGER):
            cli_main._apply_quality_auto_enable(ns)
        assert not [r for r in caplog.records if r.getMessage().startswith("quality_auto_enabled")]


class TestUnmatchedExceptionWarning:
    def _events(self, caplog: pytest.LogCaptureFixture) -> list[logging.LogRecord]:
        return [r for r in caplog.records if r.getMessage().startswith("quality_exception_unmatched")]

    def test_typo_warns(self, caplog: pytest.LogCaptureFixture) -> None:
        rules = QualityRules(
            exceptions=(
                QualityException(rule_id="dq.stale_keyword.v1", component_id="evar1", reason="ok"),
                QualityException(rule_id="dq.stale_keyword.v1", component_id="evr2", reason="typo"),
            ),
        )
        doc = _doc([_dim("evar1", "legacy_a"), _dim("evar2", "legacy_b")])
        with caplog.at_level(logging.WARNING, logger=_QUALITY_LOGGER):
            run_audits(doc, audit_naming_enabled=False, flag_stale_enabled=True, rules=rules, rsid="rs1")
        (event,) = self._events(caplog)
        assert event.levelno == logging.WARNING
        assert event.rsid == "rs1"
        assert event.rule_id == "dq.stale_keyword.v1"
        assert event.component_id == "evr2"
        assert event.component_type is None

    def test_rule_whose_audit_did_not_run_is_silent(self, caplog: pytest.LogCaptureFixture) -> None:
        rules = QualityRules(
            exceptions=(QualityException(rule_id="dq.naming_convention.v1", component_type="metrics", reason="r"),),
        )
        with caplog.at_level(logging.WARNING, logger=_QUALITY_LOGGER):
            run_audits(
                _doc([_dim("evar1", "legacy_a")]), audit_naming_enabled=False, flag_stale_enabled=True, rules=rules
            )
        assert self._events(caplog) == []

    def test_partial_cover_counts_as_matched(self, caplog: pytest.LogCaptureFixture) -> None:
        rules = QualityRules(
            exceptions=(QualityException(rule_id="dq.naming_case_inconsistency.v1", component_id="evar3", reason="r"),),
        )
        doc = _doc([_dim("evar1", "page_name"), _dim("evar2", "site_name"), _dim("evar3", "pageType")])
        with caplog.at_level(logging.WARNING, logger=_QUALITY_LOGGER):
            run_audits(doc, audit_naming_enabled=True, flag_stale_enabled=False, rules=rules)
        assert self._events(caplog) == []


class TestQualityReportCsvNeutralized:
    @pytest.mark.parametrize("lead", ["=", "+", "-", "@"])
    def test_formula_cells_neutralized(self, tmp_path: Path, lead: str) -> None:
        payload = f'{lead}HYPERLINK("http://x")'
        issue = Issue(
            SeverityLevel.LOW,
            "stale",
            "version_suffix",
            payload,
            payload,
            payload,
            {},
            component_ids=(payload,),
            excepted=True,
            exception_reason=payload,
        )
        target = tmp_path / "q.csv"
        write_quality_report(issues=[issue], summary={}, target=target, fmt="csv")
        with target.open(newline="") as fh:
            row = list(_csv.reader(fh))[1]
        for idx in (3, 4, 5, 7, 10):
            assert row[idx] == "'" + payload, (idx, row[idx])


class TestTitleCaseSmallWords:
    @pytest.mark.parametrize(
        "name",
        [
            "Time Spent per Visit",
            "Page Views per Visit",
            "Visits by Day of Week",
            "Revenue vs Target",
            "Sign in with Apple",
        ],
    )
    def test_small_words_allowed_after_first(self, name: str) -> None:
        assert matches_case_style(name, "Title Case")

    @pytest.mark.parametrize("name", ["per Visit", "the Page", "Page name", "Time Spent Per visit"])
    def test_still_rejected(self, name: str) -> None:
        assert not matches_case_style(name, "Title Case")


class TestCappedComponentIds:
    def test_human_cell_caps_at_50(self) -> None:
        ids = [f"evar{i:03d}" for i in range(120)]
        cells = dict(
            zip(
                QUALITY_ISSUE_COLUMNS,
                quality_issue_cells({"type": "case_inconsistency", "component_ids": ids, "occurrences": 120}),
                strict=True,
            )
        )
        shown = cells["component_ids"]
        assert shown.startswith("evar000, evar001")
        assert "evar049" in shown
        assert "evar050" not in shown
        assert shown.endswith("… (+70 more)")

    def test_short_list_untouched(self) -> None:
        cells = quality_issue_cells({"type": "x", "component_ids": ["a", "b"]})
        assert cells[QUALITY_ISSUE_COLUMNS.index("component_ids")] == "a, b"

    def test_json_keeps_full_list(self) -> None:
        dims = [_dim(f"evar{i}", f"page_name{i}") for i in range(60)] + [
            _dim(f"prop{i}", f"pageName{i}") for i in range(55)
        ]
        block = run_audits(_doc(dims), audit_naming_enabled=True, flag_stale_enabled=False)
        (issue,) = block["issues"]
        assert len(issue["component_ids"]) == 55


class TestCaseInconsistencyMessage:
    def test_lists_every_minority_style(self) -> None:
        doc = _doc(
            [
                _dim("evar1", "page_name"),
                _dim("evar2", "site_name"),
                _dim("evar3", "user_type"),
                _dim("evar4", "pageType"),
                _dim("evar5", "PageKind"),
                _dim("evar6", "PageMode"),
            ],
        )
        block = run_audits(doc, audit_naming_enabled=True, flag_stale_enabled=False)
        (issue,) = block["issues"]
        assert "snake_case: 3" in issue["issue"]
        assert "PascalCase: 2" in issue["issue"]
        assert "camelCase: 1" in issue["issue"]
        assert issue["occurrences"] == 3
        assert issue["component_ids"] == ["evar4", "evar5", "evar6"]


class TestExcelSummaryLayout:
    def _summary(self, tmp_path: Path, doc: SdrDocument) -> list[tuple]:
        target = tmp_path / "out.xlsx"
        ExcelWriter().write(doc, target)
        wb = load_workbook(target, read_only=True)
        try:
            rows = [tuple(r) for r in wb["Quality"].iter_rows(values_only=True)]
        finally:
            wb.close()
        header_idx = next(i for i, r in enumerate(rows) if r[0] == "severity")
        return [r[:2] for r in rows[:header_idx]]

    def test_no_exception_layout_matches_previous(self, tmp_path: Path) -> None:
        doc = _doc([_dim("evar1", "legacy_page")])
        doc = replace(doc, quality=run_audits(doc, audit_naming_enabled=False, flag_stale_enabled=True))
        fields = [r[0] for r in self._summary(tmp_path, doc)]
        assert fields == [
            "Field",
            "Verdict",
            "Policy threshold",
            "Total issues",
            "CRITICAL",
            "HIGH",
            "MEDIUM",
            "LOW",
            "INFO",
            None,
        ]

    def test_excepted_row_follows_severity_rows(self, tmp_path: Path) -> None:
        rules = QualityRules(
            exceptions=(QualityException(rule_id="dq.stale_keyword.v1", component_id="evar1", reason="r"),),
        )
        doc = _doc([_dim("evar1", "legacy_page")])
        doc = replace(doc, quality=run_audits(doc, audit_naming_enabled=False, flag_stale_enabled=True, rules=rules))
        rows = self._summary(tmp_path, doc)
        fields = [r[0] for r in rows]
        assert fields[1:9] == [
            "Verdict",
            "Policy threshold",
            "Total issues",
            "CRITICAL",
            "HIGH",
            "MEDIUM",
            "LOW",
            "INFO",
        ]
        assert rows[9] == ("Excepted", 1)
