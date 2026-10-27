"""--render-from: rebuild reports offline from saved SDR JSON or snapshots.

No client, no auth, no AA calls. Every test runs without credentials in the
environment, so any accidental auth attempt fails loudly (CONFIG, exit 10)."""

from __future__ import annotations

import csv
import dataclasses
import json
import re
import subprocess
from datetime import UTC, datetime
from pathlib import Path
from unittest.mock import MagicMock

import pytest

from aa_auto_sdr.cli.main import run
from aa_auto_sdr.output import registry
from aa_auto_sdr.sdr.document import SdrDocument
from aa_auto_sdr.snapshot.schema import document_to_envelope

FIXTURE = Path(__file__).parent.parent / "fixtures" / "sample_rs.json"
V5_FIXTURE = Path(__file__).parent.parent / "fixtures" / "snapshot_v5_metrics_only.json"
CAPTURED = datetime(2026, 4, 25, 17, 30, tzinfo=UTC)


@pytest.fixture(autouse=True)
def _no_creds(monkeypatch: pytest.MonkeyPatch, tmp_path: Path) -> None:
    for var in ("ORG_ID", "CLIENT_ID", "SECRET", "SCOPES", "AA_PROFILE", "NOTION_TOKEN"):
        monkeypatch.delenv(var, raising=False)
    monkeypatch.setenv("HOME", str(tmp_path / "home"))
    monkeypatch.chdir(tmp_path)


def _rc(argv: list[str]) -> int:
    try:
        return run(argv)
    except SystemExit as exc:  # argparse-level rejections
        return int(exc.code)


def _built_doc() -> SdrDocument:
    """A real SdrDocument from the builder (fixture data + both audits)."""
    import pandas as pd

    from aa_auto_sdr.api.client import AaClient
    from aa_auto_sdr.sdr.builder import build_sdr

    raw = json.loads(FIXTURE.read_text())
    handle = MagicMock()
    handle.getReportSuites.return_value = pd.DataFrame([raw["report_suite"]])
    handle.getDimensions.return_value = pd.DataFrame(raw["dimensions"])
    handle.getMetrics.return_value = pd.DataFrame(raw["metrics"])
    handle.getSegments.return_value = pd.DataFrame(raw["segments"])
    handle.getCalculatedMetrics.return_value = pd.DataFrame(raw["calculated_metrics"])
    handle.getVirtualReportSuites.return_value = pd.DataFrame(raw["virtual_report_suites"])
    handle.getClassificationDatasets.return_value = pd.DataFrame(raw["classification_datasets"])
    client = AaClient(handle=handle, company_id="testco")
    return build_sdr(
        client, "demo.prod", captured_at=CAPTURED, tool_version="1.21.0", audit_naming=True, flag_stale=True
    )


def _write_sdr_json(doc: SdrDocument, path: Path) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(doc.to_dict(), indent=2, sort_keys=True))
    return path


def _write_envelope(env: dict, path: Path) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(env, indent=2, sort_keys=True))
    return path


def _v1_envelope() -> dict:
    return {
        "schema": "aa-sdr-snapshot/v1",
        "rsid": "rs1",
        "captured_at": "2026-04-26T17:29:01+00:00",
        "tool_version": "1.0.0",
        "components": {
            "report_suite": {"rsid": "rs1", "name": "RS One", "timezone": None, "currency": None, "parent_rsid": None},
            "dimensions": [],
            "metrics": [],
            "segments": [],
            "calculated_metrics": [],
            "virtual_report_suites": [],
            "classifications": [],
        },
    }


def _write_live_outputs(doc: SdrDocument, out: Path) -> None:
    registry.bootstrap()
    out.mkdir(parents=True, exist_ok=True)
    for fmt in registry.resolve_formats("all"):
        writer = registry.get_writer(fmt)
        writer.write(doc, out / f"{doc.report_suite.rsid}{writer.extension}")


# --- round trip ----------------------------------------------------------------


def test_round_trip_matches_live_output_in_every_local_format(tmp_path: Path, capsys) -> None:
    doc = _built_doc()
    live = tmp_path / "live"
    _write_live_outputs(doc, live)
    saved = _write_sdr_json(doc, tmp_path / "saved" / "demo.prod.json")

    rendered = tmp_path / "rendered"
    rc = _rc(["--render-from", str(saved), "--format", "all", "--output-dir", str(rendered)])
    assert rc == 0, capsys.readouterr().err
    out = capsys.readouterr().out
    assert f"wrote: {rendered / 'demo.prod.xlsx'}" in out

    # JSON: identical apart from the two additive render keys.
    live_json = json.loads((live / "demo.prod.json").read_text())
    rendered_json = json.loads((rendered / "demo.prod.json").read_text())
    assert rendered_json.pop("rendered_by_tool_version")
    datetime.fromisoformat(rendered_json.pop("rendered_at"))
    assert rendered_json == live_json
    assert rendered_json["captured_at"] == CAPTURED.isoformat()
    assert rendered_json["tool_version"] == "1.21.0"
    assert rendered_json["quality"] is not None

    # CSV: every non-summary file is byte-identical; the summary only gains render rows.
    for live_file in sorted(live.glob("demo.prod.*.csv")):
        rendered_file = rendered / live_file.name
        if live_file.name.endswith(".summary.csv"):
            live_rows = list(csv.reader(live_file.read_text(encoding="utf-8-sig").splitlines()))
            rendered_rows = list(csv.reader(rendered_file.read_text(encoding="utf-8-sig").splitlines()))
            assert [r for r in rendered_rows if not r[0].startswith("Rendered")] == live_rows
        else:
            assert rendered_file.read_bytes() == live_file.read_bytes(), live_file.name

    # Markdown: identical once the render note and render summary rows are dropped.
    def _md(path: Path) -> list[str]:
        return [ln for ln in path.read_text(encoding="utf-8").splitlines() if ln.strip() and "Rendered" not in ln]

    assert _md(rendered / "demo.prod.md") == _md(live / "demo.prod.md")

    # HTML: identical once the render note paragraph and summary rows are dropped.
    def _html(path: Path) -> str:
        text = path.read_text(encoding="utf-8")
        text = re.sub(r'<p class="meta">Rendered [^<]*</p>', "", text)
        return re.sub(r'<tr><td class="long" title="Rendered[^"]*">.*?</tr>', "", text)

    assert _html(rendered / "demo.prod.html") == _html(live / "demo.prod.html")

    # Excel: every sheet equal; Summary only gains the render rows.
    import pandas as pd

    live_book = pd.read_excel(live / "demo.prod.xlsx", sheet_name=None)
    rendered_book = pd.read_excel(rendered / "demo.prod.xlsx", sheet_name=None)
    assert list(rendered_book) == list(live_book)
    for name, live_df in live_book.items():
        rendered_df = rendered_book[name]
        if name == "Summary":
            rendered_df = rendered_df[~rendered_df["Field"].astype(str).str.startswith("Rendered")]
            rendered_df = rendered_df.reset_index(drop=True)
        pd.testing.assert_frame_equal(rendered_df, live_df, check_dtype=False)


def test_default_format_is_excel(tmp_path: Path) -> None:
    saved = _write_sdr_json(_built_doc(), tmp_path / "demo.prod.json")
    out = tmp_path / "out"
    assert _rc(["--render-from", str(saved), "--output-dir", str(out)]) == 0
    assert [p.name for p in out.iterdir()] == ["demo.prod.xlsx"]


def test_output_dash_pipes_json_to_stdout(tmp_path: Path, capsys) -> None:
    doc = _built_doc()
    saved = _write_sdr_json(doc, tmp_path / "demo.prod.json")
    capsys.readouterr()
    assert _rc(["--render-from", str(saved), "--format", "json", "--output", "-"]) == 0
    captured = capsys.readouterr()
    payload = json.loads(captured.out)
    payload.pop("rendered_at")
    payload.pop("rendered_by_tool_version")
    assert payload == json.loads(json.dumps(doc.to_dict()))
    assert not list(tmp_path.glob("demo.prod.*.json"))


def test_agent_mode_pipes_json_to_stdout(tmp_path: Path, capsys) -> None:
    saved = _write_sdr_json(_built_doc(), tmp_path / "in" / "demo.prod.json")
    capsys.readouterr()
    assert _rc(["--render-from", str(saved), "--agent-mode"]) == 0
    assert json.loads(capsys.readouterr().out)["report_suite"]["rsid"] == "demo.prod"


def test_agent_mode_with_file_format_writes_files(tmp_path: Path, capsys) -> None:
    saved = _write_sdr_json(_built_doc(), tmp_path / "in" / "demo.prod.json")
    out = tmp_path / "out"
    assert _rc(["--render-from", str(saved), "--agent-mode", "--format", "markdown", "--output-dir", str(out)]) == 0
    assert (out / "demo.prod.md").exists()


def test_output_dash_with_non_json_format_is_output_error(tmp_path: Path, capsys) -> None:
    saved = _write_sdr_json(_built_doc(), tmp_path / "demo.prod.json")
    capsys.readouterr()
    assert _rc(["--render-from", str(saved), "--format", "markdown", "--output", "-"]) == 15
    captured = capsys.readouterr()
    assert captured.out == ""
    assert json.loads(captured.err.strip().splitlines()[-1])["error"]


def test_unknown_format_is_output_error(tmp_path: Path) -> None:
    saved = _write_sdr_json(_built_doc(), tmp_path / "demo.prod.json")
    assert _rc(["--render-from", str(saved), "--format", "console"]) == 15


# --- snapshot inputs -----------------------------------------------------------


def _coverage_states(csv_path: Path) -> dict[str, str]:
    rows = list(csv.DictReader(csv_path.read_text(encoding="utf-8-sig").splitlines()))
    return {r["family"]: r["state"] for r in rows}


def test_v1_snapshot_keeps_unknown_scope(tmp_path: Path) -> None:
    snap = _write_envelope(_v1_envelope(), tmp_path / "rs1.json")
    out = tmp_path / "out"
    assert _rc(["--render-from", str(snap), "--format", "data", "--output-dir", str(out)]) == 0
    states = _coverage_states(out / "rs1.coverage.csv")
    assert set(states.values()) == {"unknown_scope"}
    payload = json.loads((out / "rs1.json").read_text())
    assert payload["collected_components"] is None
    assert payload["captured_at"] == "2026-04-26T17:29:01+00:00"
    assert payload["tool_version"] == "1.0.0"


def test_v1_snapshot_markdown_never_says_collected_empty(tmp_path: Path) -> None:
    snap = _write_envelope(_v1_envelope(), tmp_path / "rs1.json")
    out = tmp_path / "out"
    assert _rc(["--render-from", str(snap), "--format", "markdown", "--output-dir", str(out)]) == 0
    text = (out / "rs1.md").read_text(encoding="utf-8")
    assert "Scope unknown" in text
    assert "Collected — no components" not in text


def test_v4_snapshot_keeps_fetch_status_and_quality(tmp_path: Path) -> None:
    from aa_auto_sdr.sdr.document import FetchOutcomeMeta

    doc = dataclasses.replace(
        _built_doc(),
        classifications=[],
        fetch_status={"classifications": FetchOutcomeMeta(status="degraded", expansion_level=None)},
    )
    env = document_to_envelope(doc)
    assert env["schema"] == "aa-sdr-snapshot/v4"
    snap = _write_envelope(env, tmp_path / "snap.json")
    out = tmp_path / "out"
    assert _rc(["--render-from", str(snap), "--format", "data", "--output-dir", str(out)]) == 0
    assert _coverage_states(out / "demo.prod.coverage.csv")["classifications"] == "degraded"
    payload = json.loads((out / "demo.prod.json").read_text())
    assert payload["fetch_status"] == {"classifications": {"status": "degraded", "expansion_level": None}}
    assert payload["quality"] == json.loads(json.dumps(doc.quality))


def test_v5_snapshot_keeps_filtered_scope(tmp_path: Path) -> None:
    out = tmp_path / "out"
    assert _rc(["--render-from", str(V5_FIXTURE), "--format", "csv", "--output-dir", str(out)]) == 0
    states = _coverage_states(out / "demo.prod.coverage.csv")
    assert states["dimensions"] == "excluded"
    assert states["metrics"] in ("collected", "collected_empty")


def test_rsid_at_latest_with_snapshot_dir(tmp_path: Path) -> None:
    snap_dir = tmp_path / "snaps"
    older = _v1_envelope()
    newer = _v1_envelope()
    newer["captured_at"] = "2026-05-01T00:00:00+00:00"
    _write_envelope(older, snap_dir / "rs1" / "2026-04-26T17-29-01+00-00.json")
    _write_envelope(newer, snap_dir / "rs1" / "2026-05-01T00-00-00+00-00.json")
    out = tmp_path / "out"
    argv = ["--render-from", "rs1@latest", "--snapshot-dir", str(snap_dir), "--format", "json"]
    assert _rc([*argv, "--output-dir", str(out)]) == 0
    assert json.loads((out / "rs1.json").read_text())["captured_at"] == "2026-05-01T00:00:00+00:00"
    argv[1] = "rs1@previous"
    assert _rc([*argv, "--output-dir", str(out)]) == 0
    assert json.loads((out / "rs1.json").read_text())["captured_at"] == "2026-04-26T17:29:01+00:00"


def test_rsid_at_latest_with_profile(tmp_path: Path) -> None:
    snap_dir = tmp_path / "home" / ".aa" / "orgs" / "acme" / "snapshots"
    _write_envelope(_v1_envelope(), snap_dir / "rs1" / "2026-04-26T17-29-01+00-00.json")
    out = tmp_path / "out"
    rc = _rc(["--render-from", "rs1@latest", "--profile", "acme", "--format", "json", "--output-dir", str(out)])
    assert rc == 0
    assert (out / "rs1.json").exists()


def test_rsid_at_latest_without_profile_or_dir_is_snapshot_error(tmp_path: Path, capsys) -> None:
    assert _rc(["--render-from", "rs1@latest", "--format", "json"]) == 16
    assert "requires --profile or --snapshot-dir" in capsys.readouterr().err


def test_git_token(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    repo = tmp_path / "repo"
    repo.mkdir()
    for cmd in (
        ["git", "init", "-q", "-b", "main"],
        ["git", "config", "user.email", "t@t"],
        ["git", "config", "user.name", "t"],
        ["git", "config", "commit.gpgsign", "false"],
    ):
        subprocess.run(cmd, cwd=repo, check=True)
    _write_envelope(_v1_envelope(), repo / "snaps" / "rs1.json")
    subprocess.run(["git", "add", "."], cwd=repo, check=True)
    subprocess.run(["git", "commit", "-q", "-m", "snap"], cwd=repo, check=True)
    (repo / "snaps" / "rs1.json").unlink()  # prove the content comes from git, not the worktree

    monkeypatch.chdir(repo)
    out = tmp_path / "out"
    assert _rc(["--render-from", "git:HEAD:snaps/rs1.json", "--format", "json", "--output-dir", str(out)]) == 0
    assert json.loads((out / "rs1.json").read_text())["tool_version"] == "1.0.0"


def test_template_fill_from_snapshot(tmp_path: Path) -> None:
    from openpyxl import load_workbook

    from tests.output.writers.conftest import _build_template

    template = _build_template(tmp_path / "template.xlsx")
    snap = _write_envelope(document_to_envelope(_built_doc()), tmp_path / "snap.json")
    out = tmp_path / "out"
    rc = _rc(
        [
            "--render-from",
            str(snap),
            "--template",
            str(template),
            "--template-organization",
            "Acme Corp",
            "--output-dir",
            str(out),
        ]
    )
    assert rc == 0
    wb = load_workbook(out / "demo.prod.xlsx")
    assert wb["Glossary"]["C2"].value == "Acme Corp"
    ws = wb["eVars"]
    names = {ws.cell(row=r, column=3).value: ws.cell(row=r, column=4).value for r in range(7, ws.max_row + 1)}
    assert names["evar1"] == "User ID"  # pre-seeded row filled from the snapshot
    # The template on disk is untouched.
    assert load_workbook(template)["Glossary"]["C2"].value == "Customer Org Placeholder"


def test_template_with_json_format_is_usage(tmp_path: Path, capsys) -> None:
    from tests.output.writers.conftest import _build_template

    template = _build_template(tmp_path / "template.xlsx")
    saved = _write_sdr_json(_built_doc(), tmp_path / "demo.prod.json")
    assert _rc(["--render-from", str(saved), "--template", str(template), "--format", "json"]) == 2
    assert "--template requires --format excel" in capsys.readouterr().err


# --- input errors ----------------------------------------------------------------


def test_missing_file_is_snapshot_error(tmp_path: Path, capsys) -> None:
    assert _rc(["--render-from", str(tmp_path / "nope.json")]) == 16
    assert "not found" in capsys.readouterr().err


def test_directory_is_snapshot_error(tmp_path: Path, capsys) -> None:
    assert _rc(["--render-from", str(tmp_path)]) == 16


def test_invalid_json_is_snapshot_error(tmp_path: Path, capsys) -> None:
    bad = tmp_path / "bad.json"
    bad.write_text("{not json")
    assert _rc(["--render-from", str(bad)]) == 16
    assert "not valid JSON" in capsys.readouterr().err


def test_unknown_shape_is_generic(tmp_path: Path, capsys) -> None:
    bad = _write_envelope({"hello": "world"}, tmp_path / "x.json")
    assert _rc(["--render-from", str(bad)]) == 1
    assert "unrecognized payload shape" in capsys.readouterr().err


def test_sdr_json_duplicate_ids_is_generic(tmp_path: Path, capsys) -> None:
    payload = json.loads(json.dumps(_built_doc().to_dict()))
    payload["metrics"].append(payload["metrics"][0])
    bad = _write_envelope(payload, tmp_path / "dup.json")
    out = tmp_path / "out"
    assert _rc(["--render-from", str(bad), "--output-dir", str(out)]) == 1
    err = capsys.readouterr().err
    assert "duplicate component id in metrics" in err
    assert not out.exists()


def test_sdr_json_missing_id_is_generic(tmp_path: Path, capsys) -> None:
    payload = json.loads(json.dumps(_built_doc().to_dict()))
    del payload["dimensions"][1]["id"]
    bad = _write_envelope(payload, tmp_path / "x.json")
    assert _rc(["--render-from", str(bad)]) == 1
    assert "dimensions[1]" in capsys.readouterr().err


def test_envelope_malformed_row_is_snapshot_error(tmp_path: Path, capsys) -> None:
    env = document_to_envelope(_built_doc())
    env["components"]["segments"][0]["definition"] = "not-an-object"
    bad = _write_envelope(env, tmp_path / "x.json")
    assert _rc(["--render-from", str(bad)]) == 16
    assert "segments[0]" in capsys.readouterr().err


def test_envelope_duplicate_ids_is_snapshot_error(tmp_path: Path, capsys) -> None:
    env = document_to_envelope(_built_doc())
    env["components"]["dimensions"].append(env["components"]["dimensions"][0])
    bad = _write_envelope(env, tmp_path / "x.json")
    assert _rc(["--render-from", str(bad)]) == 16
    assert "duplicate component id in dimensions" in capsys.readouterr().err


def test_unsupported_schema_is_snapshot_error(tmp_path: Path) -> None:
    env = _v1_envelope()
    env["schema"] = "aa-sdr-snapshot/v9"
    bad = _write_envelope(env, tmp_path / "x.json")
    assert _rc(["--render-from", str(bad)]) == 16


def test_pipe_errors_emit_json_envelope_and_keep_stdout_empty(tmp_path: Path, capsys) -> None:
    payload = json.loads(json.dumps(_built_doc().to_dict()))
    payload["metrics"].append(payload["metrics"][0])
    bad = _write_envelope(payload, tmp_path / "dup.json")
    capsys.readouterr()
    assert _rc(["--render-from", str(bad), "--format", "json", "--output", "-"]) == 1
    captured = capsys.readouterr()
    assert captured.out == ""
    envelope = json.loads(captured.err.strip().splitlines()[-1])
    assert envelope["error"]["type"] == "SdrPayloadError"
    assert envelope["error"]["code"] == 1


def test_pipe_snapshot_error_emits_json_envelope(tmp_path: Path, capsys) -> None:
    capsys.readouterr()
    assert _rc(["--render-from", str(tmp_path / "nope.json"), "--agent-mode"]) == 16
    captured = capsys.readouterr()
    assert captured.out == ""
    assert json.loads(captured.err.strip().splitlines()[-1])["error"]["type"] == "SnapshotResolveError"


def test_unwritable_output_dir_is_output_error(tmp_path: Path, capsys) -> None:
    saved = _write_sdr_json(_built_doc(), tmp_path / "demo.prod.json")
    blocker = tmp_path / "blocker"
    blocker.write_text("a file, not a dir")
    assert _rc(["--render-from", str(saved), "--format", "json", "--output-dir", str(blocker / "sub")]) == 15


# --- mutual exclusions -------------------------------------------------------------


def test_notion_format_points_to_push_to_notion(tmp_path: Path, capsys) -> None:
    saved = _write_sdr_json(_built_doc(), tmp_path / "demo.prod.json")
    assert _rc(["--render-from", str(saved), "--format", "notion"]) == 2
    assert "--push-to-notion" in capsys.readouterr().err


@pytest.mark.parametrize(
    "extra",
    [
        ["RS1"],
        ["--batch", "RS1"],
        ["--diff", "a.json", "b.json"],
        ["--watch", "--interval", "1h"],
        ["--list-reportsuites"],
        ["--compare-with-prev"],
        ["--trending-window", "30d"],
        ["--stats"],
        ["--snapshot"],
        ["--auto-snapshot"],
        ["--auto-prune", "--keep-last", "3"],
        ["--git-commit"],
        ["--metrics-only"],
        ["--dimensions-only"],
        ["--dry-run"],
        ["--audit-naming"],
        ["--flag-stale"],
        ["--quality-report", "json"],
        ["--fail-on-quality", "HIGH"],
        ["--quality-policy", "policy.json"],
        ["--name-match", "exact"],
        ["--workers", "2"],
        ["--enable-cache"],
        ["--max-retries", "5"],
        ["--filter", "x"],
        ["--limit", "3"],
        ["--run-summary-json", "summary.json"],
        ["--show-timings"],
        ["--open"],
        ["--push-to-notion", "x.json"],
        ["--notion-force-new"],
        ["--notion-create-database"],
        ["--ignore-fields", "description"],
        ["--warn-threshold", "3"],
    ],
)
def test_generation_and_other_mode_flags_are_rejected(extra: list[str], tmp_path: Path, capsys) -> None:
    saved = _write_sdr_json(_built_doc(), tmp_path / "demo.prod.json")
    out = tmp_path / "out"
    assert _rc(["--render-from", str(saved), *extra, "--output-dir", str(out)]) == 2
    assert "--render-from" in capsys.readouterr().err
    assert not out.exists()


def test_non_utf8_file_is_snapshot_error(tmp_path: Path, capsys) -> None:
    bad = tmp_path / "latin1.json"
    bad.write_bytes(b'{"name": "\xff\xfe"}')
    assert _rc(["--render-from", str(bad)]) == 16
    assert "could not read" in capsys.readouterr().err
