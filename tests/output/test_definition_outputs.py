"""Definition summaries, reference columns, and the references appendix in every report format."""

from __future__ import annotations

import csv as _csv
import json
from datetime import UTC, datetime
from pathlib import Path
from typing import Any

import pytest

from aa_auto_sdr.api import models
from aa_auto_sdr.sdr.document import SdrDocument

RATE_DEF = {
    "func": "calc-metric",
    "formula": {
        "func": "divide",
        "col1": {"func": "metric", "name": "metrics/orders"},
        "col2": {"func": "metric", "name": "metrics/visits"},
    },
}
MOBILE_CM_DEF = {
    "func": "calc-metric",
    "formula": {"func": "segment", "segment_id": "s_mobile", "metric": {"func": "metric", "name": "cm_rate"}},
}
MOBILE_SEG_DEF = {
    "func": "segment",
    "container": {
        "func": "container",
        "context": "hits",
        "pred": {"func": "streq", "str": "Phone", "val": {"func": "attr", "name": "variables/mobiledevicetype"}},
    },
}


def _doc(*, empty: bool = False) -> SdrDocument:
    metrics = [
        models.Metric(id=mid, name=name, type="int", category=None, precision=0, segmentable=True, description=None)
        for mid, name in (("metrics/orders", "Orders"), ("metrics/visits", "Visits"))
    ]

    def cm(cid: str, name: str, definition: dict[str, Any]) -> models.CalculatedMetric:
        return models.CalculatedMetric(
            id=cid,
            name=name,
            description=None,
            rsid="rs",
            owner_id=None,
            polarity="positive",
            precision=2,
            type="decimal",
            definition=definition,
        )

    return SdrDocument(
        report_suite=models.ReportSuite(rsid="rs", name="RS", timezone=None, currency=None, parent_rsid=None),
        dimensions=[],
        metrics=[] if empty else metrics,
        segments=[]
        if empty
        else [
            models.Segment(
                id="s_mobile", name="Mobile", description=None, rsid="rs", owner_id=None, definition=MOBILE_SEG_DEF
            )
        ],
        calculated_metrics=[]
        if empty
        else [cm("cm_rate", "Rate", RATE_DEF), cm("cm_mobile", "Mobile Rate", MOBILE_CM_DEF)],
        virtual_report_suites=[]
        if empty
        else [
            models.VirtualReportSuite(
                id="vrs_m",
                name="Mobile VRS",
                parent_rsid="rs",
                timezone=None,
                description=None,
                segment_list=["s_mobile", "s_gone"],
            )
        ],
        classifications=[],
        captured_at=datetime(2026, 1, 1, tzinfo=UTC),
        tool_version="1.0.0",
        collected_components=(
            "calculated_metrics",
            "classifications",
            "dimensions",
            "metrics",
            "segments",
            "virtual_report_suites",
        ),
    )


def _read_csv(path: Path) -> list[dict[str, str]]:
    with path.open(encoding="utf-8-sig") as fh:
        return list(_csv.DictReader(fh))


class TestCsv:
    def test_component_columns_are_additive(self, tmp_path: Path) -> None:
        from aa_auto_sdr.output.writers.csv import CsvWriter

        CsvWriter().write(_doc(), tmp_path / "out.csv")
        cms = _read_csv(tmp_path / "out.calculated_metrics.csv")
        header = list(cms[0])
        assert header[-3:] == ["formula_summary", "metric_references", "segment_references"]
        assert "definition" in header
        rate = next(r for r in cms if r["id"] == "cm_rate")
        assert rate["formula_summary"] == "metrics/orders / metrics/visits"
        assert rate["metric_references"] == "metrics/orders, metrics/visits"
        assert rate["segment_references"] == ""
        assert json.loads(rate["definition"]) == RATE_DEF
        segs = _read_csv(tmp_path / "out.segments.csv")
        assert list(segs[0])[-3:] == ["definition_summary", "metric_references", "segment_references"]
        assert segs[0]["definition_summary"] == "Hit where variables/mobiledevicetype = 'Phone'"

    def test_references_file(self, tmp_path: Path) -> None:
        from aa_auto_sdr.output.writers.csv import CsvWriter

        paths = CsvWriter().write(_doc(), tmp_path / "out.csv")
        ref_path = tmp_path / "out.references.csv"
        assert ref_path in paths
        rows = _read_csv(ref_path)
        assert list(rows[0]) == [
            "family",
            "id",
            "name",
            "relation",
            "kind",
            "target_family",
            "target_id",
            "target_name",
            "status",
        ]
        gone = next(r for r in rows if r["target_id"] == "s_gone")
        assert gone["status"] == "unresolved"
        assert gone["target_family"] == ""
        by = [r for r in rows if r["id"] == "s_mobile" and r["relation"] == "referenced_by"]
        assert [(r["target_family"], r["target_id"]) for r in by] == [
            ("calculated_metrics", "cm_mobile"),
            ("virtual_report_suites", "vrs_m"),
        ]

    def test_empty_document_writes_headers(self, tmp_path: Path) -> None:
        from aa_auto_sdr.output.writers.csv import CsvWriter

        CsvWriter().write(_doc(empty=True), tmp_path / "out.csv")
        with (tmp_path / "out.calculated_metrics.csv").open(encoding="utf-8-sig") as fh:
            header = next(_csv.reader(fh))
        assert header[-3:] == ["formula_summary", "metric_references", "segment_references"]
        with (tmp_path / "out.references.csv").open(encoding="utf-8-sig") as fh:
            assert list(_csv.reader(fh)) == [
                ["family", "id", "name", "relation", "kind", "target_family", "target_id", "target_name", "status"]
            ]


class TestExcel:
    def test_columns_and_references_sheet(self, tmp_path: Path) -> None:
        import openpyxl

        from aa_auto_sdr.output.writers.excel import ExcelWriter

        (path,) = ExcelWriter().write(_doc(), tmp_path / "out.xlsx")
        wb = openpyxl.load_workbook(path)
        assert wb.sheetnames[-1] == "References"
        assert wb.sheetnames.index("References") == wb.sheetnames.index("Classifications") + 1
        ws = wb["Calculated Metrics"]
        header = [c.value for c in ws[1]]
        assert header[-3:] == ["formula_summary", "metric_references", "segment_references"]
        values = {row[0]: row for row in ws.iter_rows(min_row=2, values_only=True)}
        assert values["cm_mobile"][-3:] == ("cm_rate[segment s_mobile]", "cm_rate", "s_mobile")
        refs = wb["References"]
        rheader = [c.value for c in refs[1]]
        assert rheader == [
            "Family",
            "ID",
            "Name",
            "Relation",
            "Kind",
            "Target Family",
            "Target ID",
            "Target Name",
            "Status",
        ]
        rows = list(refs.iter_rows(min_row=2, values_only=True))
        assert (
            "Virtual Report Suites",
            "vrs_m",
            "Mobile VRS",
            "references",
            "segment",
            None,
            "s_gone",
            None,
            "unresolved",
        ) in rows
        assert any(r[:4] == ("Calculated Metrics", "cm_rate", "Rate", "referenced by") for r in rows)

    def test_empty_references_sheet_has_header(self, tmp_path: Path) -> None:
        import openpyxl

        from aa_auto_sdr.output.writers.excel import ExcelWriter

        (path,) = ExcelWriter().write(_doc(empty=True), tmp_path / "out.xlsx")
        ws = openpyxl.load_workbook(path)["References"]
        assert next(c.value for c in ws[1]) == "Family"
        assert ws.max_row == 1


class TestMarkdown:
    def test_summary_columns_and_section(self, tmp_path: Path) -> None:
        from aa_auto_sdr.output.writers.markdown import MarkdownWriter

        (path,) = MarkdownWriter().write(_doc(), tmp_path / "out.md")
        text = path.read_text()
        assert "| formula_summary | metric_references | segment_references |" in text
        assert "| `metrics/orders / metrics/visits` | metrics/orders, metrics/visits |  |" in text
        assert "| definition_summary | metric_references | segment_references |" in text
        assert "## References\n" in text
        assert text.index("## References") > text.index("## Classifications")
        assert (
            "| Virtual Report Suites | vrs_m | Mobile VRS | references | segment |  | s_gone |  | unresolved |" in text
        )
        assert "referenced by" in text

    def test_empty_references_section(self, tmp_path: Path) -> None:
        from aa_auto_sdr.output.writers.markdown import MarkdownWriter

        (path,) = MarkdownWriter().write(_doc(empty=True), tmp_path / "out.md")
        assert "## References\n\n_No references found._\n" in path.read_text()


class TestMarkdownSpecials:
    @pytest.mark.parametrize(
        ("value", "cell"),
        [
            ("*sale*", "| `Hit where v/x = '*sale*'` |"),
            ("_x_", "| `Hit where v/x = '_x_'` |"),
            ("a\\b", "| `Hit where v/x = 'a\\\\b'` |"),
            ("a|b", "| `Hit where v/x = 'a\\|b'` |"),
            ("a`b", "| ``Hit where v/x = 'a`b'`` |"),
        ],
    )
    def test_summary_cells_are_inline_code(self, tmp_path: Path, value: str, cell: str) -> None:
        from dataclasses import replace

        from aa_auto_sdr.output.writers.markdown import MarkdownWriter

        definition = {
            "func": "segment",
            "container": {
                "func": "container",
                "context": "hits",
                "pred": {"func": "streq", "str": value, "val": {"func": "attr", "name": "v/x"}},
            },
        }
        doc = _doc()
        doc = replace(doc, segments=[replace(doc.segments[0], definition=definition)])
        (path,) = MarkdownWriter().write(doc, tmp_path / "out.md")
        line = next(ln for ln in path.read_text().splitlines() if ln.startswith("| s_mobile"))
        assert cell in line


class TestHtml:
    def test_summary_columns_and_section(self, tmp_path: Path) -> None:
        from aa_auto_sdr.output.writers.html import HtmlWriter

        (path,) = HtmlWriter().write(_doc(), tmp_path / "out.html")
        text = path.read_text()
        assert "<th>formula_summary</th><th>metric_references</th><th>segment_references</th>" in text
        assert "metrics/orders / metrics/visits" in text
        assert "<h2>References" in text
        assert text.index("<h2>References") > text.index("<h2>Classifications")
        assert "unresolved" in text
        assert "Hit where variables/mobiledevicetype = &#x27;Phone&#x27;" in text

    def test_empty_references_section(self, tmp_path: Path) -> None:
        from aa_auto_sdr.output.writers.html import HtmlWriter

        (path,) = HtmlWriter().write(_doc(empty=True), tmp_path / "out.html")
        assert "No references found." in path.read_text()


class TestJson:
    def test_fields_and_references_block(self, tmp_path: Path) -> None:
        from aa_auto_sdr.output.writers.json import JsonWriter

        (path,) = JsonWriter().write(_doc(), tmp_path / "out.json")
        payload = json.loads(path.read_text())
        rate = next(c for c in payload["calculated_metrics"] if c["id"] == "cm_rate")
        assert rate["formula_summary"] == "metrics/orders / metrics/visits"
        assert rate["metric_references"] == ["metrics/orders", "metrics/visits"]
        assert rate["segment_references"] == []
        assert rate["definition"] == RATE_DEF
        assert payload["segments"][0]["definition_summary"].startswith("Hit where")
        refs = payload["references"]
        assert set(refs) == {"calculated_metrics", "segments", "virtual_report_suites"}
        vrs = refs["virtual_report_suites"][0]
        assert vrs["references"][0] == {
            "kind": "segment",
            "id": "s_gone",
            "family": None,
            "name": None,
            "status": "unresolved",
        }

    def test_retained_snapshot_payload_stays_plain(self, tmp_path: Path) -> None:
        from aa_auto_sdr.output.writers.json import JsonWriter

        doc = _doc()
        _, payload = JsonWriter()._write_and_retain_payload(doc, tmp_path / "out.json")
        assert payload == doc.to_dict()
        assert "references" not in payload
        assert "formula_summary" not in payload["calculated_metrics"][0]

    def test_written_json_renders_back(self, tmp_path: Path) -> None:
        """--render-from accepts the enriched JSON: unknown keys are ignored."""
        from aa_auto_sdr.output.writers.json import JsonWriter
        from aa_auto_sdr.sdr.rebuild import document_from_payload

        doc = _doc()
        (path,) = JsonWriter().write(doc, tmp_path / "out.json")
        rebuilt = document_from_payload(
            json.loads(path.read_text()), rendered_at=datetime(2026, 2, 1, tzinfo=UTC), rendering_tool_version="x"
        )
        assert rebuilt.calculated_metrics == doc.calculated_metrics
        assert rebuilt.segments == doc.segments


@pytest.mark.parametrize("family", ["calculated_metrics", "segments"])
def test_snapshot_envelope_has_no_derived_fields(family: str) -> None:
    from aa_auto_sdr.snapshot.schema import document_to_envelope

    env = document_to_envelope(_doc())
    assert "references" not in env["components"]
    for row in env["components"][family]:
        assert not {"formula_summary", "definition_summary", "metric_references", "segment_references"} & set(row)
