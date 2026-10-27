"""Coverage and quality status in every local report format.

A reader of any report (Excel, CSV, Markdown, HTML) must be able to tell, for
each component family, whether it was collected with rows, collected empty,
excluded from the capture, degraded, partial (with its expansion level), or of
unknown scope; and whether a quality audit ran, with its verdict and findings.
"""

from __future__ import annotations

import csv as _csv
from dataclasses import replace
from datetime import UTC, datetime
from pathlib import Path

import pytest
from openpyxl import load_workbook

from aa_auto_sdr.api import models as m
from aa_auto_sdr.output.coverage import (
    CoverageState,
    coverage_rows,
    quality_status,
    quality_status_line,
)
from aa_auto_sdr.output.writers.csv import CsvWriter
from aa_auto_sdr.output.writers.excel import ExcelWriter
from aa_auto_sdr.output.writers.html import HtmlWriter
from aa_auto_sdr.output.writers.markdown import MarkdownWriter
from aa_auto_sdr.sdr.document import COMPONENT_FAMILIES, FetchOutcomeMeta, SdrDocument

_DIMENSION = m.Dimension(
    id="variables/evar1",
    name="User ID",
    type="string",
    category="Conversion",
    parent="",
    pathable=False,
    description=None,
)
_METRIC = m.Metric(
    id="metrics/pageviews",
    name="Page Views",
    type="int",
    category=None,
    precision=0,
    segmentable=True,
    description=None,
)
_CLASSIFICATION = m.ClassificationDataset(id="ds_1", name="Campaign <Codes>", rsid="demo.prod")

_AUDITED_QUALITY = {
    "naming_audit": {"total_components": 1, "case_styles": {}, "prefix_groups": {}, "recommendations": []},
    "stale_components": [],
    "issues": [
        {
            "severity": "HIGH",
            "category": "stale",
            "type": "dimension",
            "item_id": "variables/evar1",
            "item_name": "User ID test",
            "issue": "stale keyword: test",
            "details": {"reason": "stale_keyword:test"},
        },
        {
            "severity": "LOW",
            "category": "naming",
            "type": "dimensions",
            "item_id": "dimensions",
            "item_name": "",
            "issue": "mixed case styles",
            "details": {},
        },
    ],
    "summary": {
        "by_severity": {"CRITICAL": 0, "HIGH": 1, "MEDIUM": 0, "LOW": 1, "INFO": 0},
        "total": 2,
        "verdict": "fail",
        "policy_threshold": "HIGH",
    },
}


def _doc(**overrides) -> SdrDocument:
    base = SdrDocument(
        report_suite=m.ReportSuite(rsid="demo.prod", name="Demo", timezone=None, currency=None, parent_rsid=None),
        dimensions=[],
        metrics=[],
        segments=[],
        calculated_metrics=[],
        virtual_report_suites=[],
        classifications=[],
        captured_at=datetime(2026, 4, 25, tzinfo=UTC),
        tool_version="9.9.9",
        collected_components=COMPONENT_FAMILIES,
    )
    return replace(base, **overrides)


@pytest.fixture
def mixed_doc() -> SdrDocument:
    """Full scope: dimensions healthy with rows, metrics/segments/CMs healthy
    empty, VRS degraded empty, classifications partial with a level. Not audited."""
    return _doc(
        dimensions=[_DIMENSION],
        classifications=[_CLASSIFICATION],
        fetch_status={
            "virtual_report_suites": FetchOutcomeMeta(status="degraded", expansion_level=None),
            "classifications": FetchOutcomeMeta(status="partial", expansion_level="rsid_only"),
        },
    )


@pytest.fixture
def metrics_only_doc() -> SdrDocument:
    """Metrics-only capture scope, audited with a failing gate."""
    return _doc(metrics=[_METRIC], collected_components=("metrics",), quality=_AUDITED_QUALITY)


@pytest.fixture
def unknown_scope_doc() -> SdrDocument:
    return _doc(collected_components=None)


# --- pure helper ------------------------------------------------------------


def test_coverage_rows_mixed_states(mixed_doc: SdrDocument) -> None:
    rows = {r.family: r for r in coverage_rows(mixed_doc)}
    assert set(rows) == set(COMPONENT_FAMILIES)
    assert rows["dimensions"].state is CoverageState.COLLECTED
    assert rows["dimensions"].count == 1
    assert rows["metrics"].state is CoverageState.COLLECTED_EMPTY
    assert rows["virtual_report_suites"].state is CoverageState.DEGRADED
    assert rows["virtual_report_suites"].count == 0
    assert rows["classifications"].state is CoverageState.PARTIAL
    assert rows["classifications"].expansion_level == "rsid_only"
    assert rows["classifications"].count == 1


def test_coverage_rows_follow_report_section_order(mixed_doc: SdrDocument) -> None:
    assert [r.family for r in coverage_rows(mixed_doc)] == [
        "dimensions",
        "metrics",
        "segments",
        "calculated_metrics",
        "virtual_report_suites",
        "classifications",
    ]


def test_coverage_rows_excluded(metrics_only_doc: SdrDocument) -> None:
    rows = {r.family: r for r in coverage_rows(metrics_only_doc)}
    assert rows["metrics"].state is CoverageState.COLLECTED
    for family in COMPONENT_FAMILIES:
        if family != "metrics":
            assert rows[family].state is CoverageState.EXCLUDED


def test_coverage_rows_unknown_scope_is_never_healthy_empty(unknown_scope_doc: SdrDocument) -> None:
    rows = coverage_rows(unknown_scope_doc)
    assert {r.state for r in rows} == {CoverageState.UNKNOWN_SCOPE}


def test_coverage_rows_unknown_scope_keeps_definite_fetch_states() -> None:
    doc = _doc(
        collected_components=None,
        fetch_status={"virtual_report_suites": FetchOutcomeMeta(status="degraded", expansion_level=None)},
    )
    rows = {r.family: r for r in coverage_rows(doc)}
    assert rows["virtual_report_suites"].state is CoverageState.DEGRADED
    assert rows["metrics"].state is CoverageState.UNKNOWN_SCOPE


def test_coverage_labels_are_distinct(mixed_doc, metrics_only_doc, unknown_scope_doc) -> None:
    labels = {}
    for doc in (mixed_doc, metrics_only_doc, unknown_scope_doc):
        for row in coverage_rows(doc):
            labels.setdefault(row.state, row.label)
    assert len(labels) == len(CoverageState)
    assert len(set(labels.values())) == len(CoverageState)
    assert labels[CoverageState.EXCLUDED] == "Not collected in this capture"
    assert labels[CoverageState.COLLECTED_EMPTY] == "Collected — no components"
    assert labels[CoverageState.DEGRADED] == "Fetch degraded — no data"
    assert "rsid_only" in labels[CoverageState.PARTIAL]


def test_quality_status_not_audited() -> None:
    status = quality_status(None)
    assert status.audited is False
    assert quality_status_line(status) == "No quality audit ran."


def test_quality_status_audited_with_verdict() -> None:
    status = quality_status(_AUDITED_QUALITY)
    assert status.audited is True
    assert status.verdict == "fail"
    assert status.threshold == "HIGH"
    assert status.total == 2
    assert list(status.by_severity) == ["CRITICAL", "HIGH", "MEDIUM", "LOW", "INFO"]
    line = quality_status_line(status)
    assert "FAIL" in line
    assert "HIGH" in line
    assert "2 issues" in line


def test_quality_status_audited_without_policy() -> None:
    quality = {**_AUDITED_QUALITY, "summary": {"by_severity": {}, "total": 0, "verdict": "n/a"}, "issues": []}
    line = quality_status_line(quality_status(quality))
    assert "no gate policy" in line
    assert "0 issues" in line


def test_quality_status_legacy_block_without_summary() -> None:
    """A pre-severity quality block (naming_audit/stale_components only) still
    reads as audited, with zero promoted issues."""
    status = quality_status({"naming_audit": {}, "stale_components": []})
    assert status.audited is True
    assert status.total == 0
    assert status.verdict is None


# --- Excel ------------------------------------------------------------------


def _sheet_rows(path: Path, name: str) -> list[tuple]:
    wb = load_workbook(path, read_only=True)
    try:
        return [tuple(c for c in row) for row in wb[name].iter_rows(values_only=True)]
    finally:
        wb.close()


def test_excel_coverage_sheet_follows_summary(mixed_doc: SdrDocument, tmp_path: Path) -> None:
    target = tmp_path / "out.xlsx"
    ExcelWriter().write(mixed_doc, target)
    wb = load_workbook(target, read_only=True)
    names = wb.sheetnames
    wb.close()
    assert names[:3] == ["Summary", "Coverage", "Quality"]
    assert names[3:] == [
        "Dimensions",
        "Metrics",
        "Segments",
        "Calculated Metrics",
        "Virtual Report Suites",
        "Classifications",
    ]
    rows = _sheet_rows(target, "Coverage")
    assert rows[0] == ("Family", "State", "Count", "Expansion Level")
    by_family = {r[0]: r for r in rows[1:]}
    assert by_family["Dimensions"][1:3] == ("Collected", 1)
    assert by_family["Metrics"][1:3] == ("Collected — no components", 0)
    assert by_family["Virtual Report Suites"][1] == "Fetch degraded — no data"
    assert by_family["Classifications"][1].startswith("Partial fetch")
    assert by_family["Classifications"][3] == "rsid_only"


def test_excel_coverage_excluded_and_unknown(metrics_only_doc, unknown_scope_doc, tmp_path: Path) -> None:
    excluded = tmp_path / "excluded.xlsx"
    ExcelWriter().write(metrics_only_doc, excluded)
    states = {r[0]: r[1] for r in _sheet_rows(excluded, "Coverage")[1:]}
    assert states["Metrics"] == "Collected"
    assert states["Dimensions"] == "Not collected in this capture"

    unknown = tmp_path / "unknown.xlsx"
    ExcelWriter().write(unknown_scope_doc, unknown)
    states = {r[0]: r[1] for r in _sheet_rows(unknown, "Coverage")[1:]}
    assert all(s.startswith("Scope unknown") for s in states.values())


def test_excel_quality_sheet_not_audited(mixed_doc: SdrDocument, tmp_path: Path) -> None:
    target = tmp_path / "out.xlsx"
    ExcelWriter().write(mixed_doc, target)
    assert _sheet_rows(target, "Quality") == [("No quality audit ran.",)]


def test_excel_quality_sheet_traces_verdict_to_findings(metrics_only_doc: SdrDocument, tmp_path: Path) -> None:
    target = tmp_path / "out.xlsx"
    ExcelWriter().write(metrics_only_doc, target)
    rows = _sheet_rows(target, "Quality")
    header_idx = next(i for i, r in enumerate(rows) if r[0] == "severity")
    summary = {r[0]: r[1] for r in rows[:header_idx] if r[0]}
    assert summary["Verdict"] == "fail"
    assert summary["Policy threshold"] == "HIGH"
    assert summary["Total issues"] == 2
    assert summary["HIGH"] == 1
    assert rows[header_idx][:6] == ("severity", "category", "type", "item_id", "item_name", "issue")
    issue_rows = [r for r in rows[header_idx + 1 :] if r[0]]
    assert len(issue_rows) == 2
    # The failing verdict is explained by a finding at the policy threshold.
    assert any(r[0] == "HIGH" and r[3] == "variables/evar1" for r in issue_rows)


def test_excel_existing_sheet_columns_unchanged(mixed_doc: SdrDocument, tmp_path: Path) -> None:
    target = tmp_path / "out.xlsx"
    ExcelWriter().write(mixed_doc, target)
    summary = _sheet_rows(target, "Summary")
    assert summary[0] == ("Field", "Value")
    assert summary[1] == ("RSID", "demo.prod")
    dims = _sheet_rows(target, "Dimensions")
    assert dims[0][0] == "id"


# --- CSV --------------------------------------------------------------------


def _read_csv(path: Path) -> list[list[str]]:
    with path.open(encoding="utf-8-sig", newline="") as fh:
        return list(_csv.reader(fh))


def test_csv_coverage_file_always_written(mixed_doc: SdrDocument, tmp_path: Path) -> None:
    paths = CsvWriter().write(mixed_doc, tmp_path / "out.csv")
    names = [p.name for p in paths]
    assert "out.coverage.csv" in names
    assert "out.quality_issues.csv" not in names
    assert not (tmp_path / "out.quality_issues.csv").exists()
    rows = _read_csv(tmp_path / "out.coverage.csv")
    assert rows[0] == ["family", "state", "count", "expansion_level"]
    by_family = {r[0]: r for r in rows[1:]}
    assert by_family["dimensions"][1:3] == ["collected", "1"]
    assert by_family["metrics"][1:3] == ["collected_empty", "0"]
    assert by_family["virtual_report_suites"][1] == "degraded"
    assert by_family["classifications"][1:] == ["partial", "1", "rsid_only"]


def test_csv_coverage_excluded_and_unknown(metrics_only_doc, unknown_scope_doc, tmp_path: Path) -> None:
    CsvWriter().write(metrics_only_doc, tmp_path / "a.csv")
    states = {r[0]: r[1] for r in _read_csv(tmp_path / "a.coverage.csv")[1:]}
    assert states["metrics"] == "collected"
    assert states["segments"] == "excluded"

    CsvWriter().write(unknown_scope_doc, tmp_path / "b.csv")
    states = {r[0]: r[1] for r in _read_csv(tmp_path / "b.coverage.csv")[1:]}
    assert set(states.values()) == {"unknown_scope"}


def test_csv_quality_issues_traces_verdict(metrics_only_doc: SdrDocument, tmp_path: Path) -> None:
    paths = CsvWriter().write(metrics_only_doc, tmp_path / "out.csv")
    assert tmp_path / "out.quality_issues.csv" in paths
    rows = _read_csv(tmp_path / "out.quality_issues.csv")
    assert rows[0][:7] == ["severity", "category", "type", "item_id", "item_name", "issue", "details"]
    assert rows[0][7:] == ["rule_id", "component_ids", "occurrences", "exception"]
    assert len(rows) == 3
    assert rows[1][0] == "HIGH"
    assert rows[1][3] == "variables/evar1"
    summary = {r[0]: r[1] for r in _read_csv(tmp_path / "out.summary.csv")[1:]}
    assert summary["Quality verdict"] == "fail"
    assert summary["Quality policy threshold"] == "HIGH"
    assert summary["Quality issues"] == "2"


def test_csv_summary_quality_not_run(mixed_doc: SdrDocument, tmp_path: Path) -> None:
    CsvWriter().write(mixed_doc, tmp_path / "out.csv")
    summary = {r[0]: r[1] for r in _read_csv(tmp_path / "out.summary.csv")[1:]}
    assert summary["Quality verdict"] == "not run"


# --- Markdown ---------------------------------------------------------------


def _markdown(doc: SdrDocument, tmp_path: Path) -> str:
    target = tmp_path / "out.md"
    MarkdownWriter().write(doc, target)
    return target.read_text(encoding="utf-8")


def test_markdown_coverage_section_after_header(mixed_doc: SdrDocument, tmp_path: Path) -> None:
    content = _markdown(mixed_doc, tmp_path)
    assert content.index("## Coverage") < content.index("## Summary")
    assert content.index("# SDR") < content.index("## Coverage")
    assert "| Family | State | Count | Expansion level |" in content
    assert "| Classifications | Partial fetch (expansion level: rsid_only) | 1 | rsid_only |" in content
    assert "No quality audit ran." in content


def test_markdown_empty_sections_show_real_state(mixed_doc, metrics_only_doc, unknown_scope_doc, tmp_path) -> None:
    content = _markdown(mixed_doc, tmp_path)
    assert "_(none)_" not in content
    metrics_section = content.split("## Metrics\n", 1)[1].split("\n## ", 1)[0]
    assert "_Collected — no components_" in metrics_section
    vrs_section = content.split("## Virtual Report Suites\n", 1)[1].split("\n## ", 1)[0]
    assert "_Fetch degraded — no data_" in vrs_section

    content = _markdown(metrics_only_doc, tmp_path)
    seg_section = content.split("## Segments\n", 1)[1].split("\n## ", 1)[0]
    assert "_Not collected in this capture_" in seg_section

    content = _markdown(unknown_scope_doc, tmp_path)
    seg_section = content.split("## Segments\n", 1)[1].split("\n## ", 1)[0]
    assert "Scope unknown" in seg_section
    assert "Collected — no components" not in content


def test_markdown_partial_section_with_rows_notes_state(mixed_doc: SdrDocument, tmp_path: Path) -> None:
    content = _markdown(mixed_doc, tmp_path)
    cls_section = content.split("## Classifications\n", 1)[1]
    assert "Partial fetch (expansion level: rsid_only)" in cls_section
    assert "ds_1" in cls_section


def test_markdown_quality_verdict(metrics_only_doc: SdrDocument, tmp_path: Path) -> None:
    content = _markdown(metrics_only_doc, tmp_path)
    assert "## Quality" in content
    quality = content.split("## Quality\n", 1)[1].split("\n## ", 1)[0]
    assert "FAIL" in quality
    assert "HIGH" in quality


# --- HTML -------------------------------------------------------------------


def _html(doc: SdrDocument, tmp_path: Path) -> str:
    target = tmp_path / "out.html"
    HtmlWriter().write(doc, target)
    return target.read_text(encoding="utf-8")


def test_html_coverage_section_after_header(mixed_doc: SdrDocument, tmp_path: Path) -> None:
    content = _html(mixed_doc, tmp_path)
    assert content.index("</header>") < content.index("<h2>Coverage</h2>") < content.index("<h2>Summary</h2>")
    assert "Fetch degraded — no data" in content
    assert "rsid_only" in content
    assert "No quality audit ran." in content
    assert "<em>(none)</em>" not in content


def test_html_empty_sections_show_real_state(metrics_only_doc, unknown_scope_doc, tmp_path) -> None:
    content = _html(metrics_only_doc, tmp_path)
    seg = content.split("<h2>Segments", 1)[1].split("</section>", 1)[0]
    assert "Not collected in this capture" in seg
    assert "FAIL" in content

    content = _html(unknown_scope_doc, tmp_path)
    seg = content.split("<h2>Segments", 1)[1].split("</section>", 1)[0]
    assert "Scope unknown" in seg


def test_html_escapes_coverage_and_quality_text(tmp_path: Path) -> None:
    quality = {
        **_AUDITED_QUALITY,
        "summary": {**_AUDITED_QUALITY["summary"], "verdict": "<b>fail</b>", "policy_threshold": "<i>HIGH</i>"},
    }
    doc = _doc(
        quality=quality,
        fetch_status={"classifications": FetchOutcomeMeta(status="partial", expansion_level="<script>x</script>")},
    )
    content = _html(doc, tmp_path)
    assert "<script>x</script>" not in content
    assert "&lt;script&gt;" in content
    assert "<b>fail</b>" not in content
    assert "<i>HIGH</i>" not in content


def test_coverage_accepts_dict_fetch_meta_and_degraded_with_rows() -> None:
    doc = _doc(
        dimensions=[_DIMENSION],
        fetch_status={"dimensions": {"status": "degraded", "expansion_level": None}},
    )
    row = coverage_rows(doc)[0]
    assert row.state is CoverageState.DEGRADED
    assert row.label == "Fetch degraded"


def test_quality_status_counts_issues_when_summary_missing() -> None:
    status = quality_status({"issues": _AUDITED_QUALITY["issues"]})
    assert status.total == 2
    assert status.by_severity == {"CRITICAL": 0, "HIGH": 1, "MEDIUM": 0, "LOW": 1, "INFO": 0}
    assert quality_status_line(status).startswith("Quality audit ran with no gate policy: 2 issues")
