"""Rule IDs, component IDs, occurrences, and exception state in every output."""

from __future__ import annotations

import csv as _csv
import json
from datetime import UTC, datetime
from pathlib import Path

from openpyxl import load_workbook

from aa_auto_sdr.api import models as m
from aa_auto_sdr.output.coverage import quality_status, quality_status_line
from aa_auto_sdr.output.writers.csv import CsvWriter
from aa_auto_sdr.output.writers.excel import ExcelWriter
from aa_auto_sdr.output.writers.html import HtmlWriter
from aa_auto_sdr.output.writers.json import JsonWriter
from aa_auto_sdr.output.writers.markdown import MarkdownWriter
from aa_auto_sdr.pipeline.single import _rehydrate_issue
from aa_auto_sdr.sdr.document import COMPONENT_FAMILIES, SdrDocument
from aa_auto_sdr.sdr.quality import NamingRule, QualityException, QualityRules, SeverityLevel, run_audits
from aa_auto_sdr.sdr.quality_policy import write_quality_report


def _dim(cid: str, name: str) -> m.Dimension:
    return m.Dimension(id=cid, name=name, type="string", category=None, parent="", pathable=False, description=None)


def _doc() -> SdrDocument:
    base = SdrDocument(
        report_suite=m.ReportSuite(rsid="demo.prod", name="Demo", timezone=None, currency=None, parent_rsid=None),
        dimensions=[_dim("evar1", "legacy_page"), _dim("evar2", "PageName")],
        metrics=[],
        segments=[],
        calculated_metrics=[],
        virtual_report_suites=[],
        classifications=[],
        captured_at=datetime(2026, 4, 25, tzinfo=UTC),
        tool_version="9.9.9",
        collected_components=COMPONENT_FAMILIES,
    )
    rules = QualityRules(
        naming={"dimensions": NamingRule(case="snake_case")},
        exceptions=(
            QualityException(rule_id="dq.stale_keyword.v1", component_id="evar1", reason="kept for audit trail"),
        ),
    )
    quality = run_audits(
        base,
        audit_naming_enabled=True,
        flag_stale_enabled=True,
        fail_on_quality=SeverityLevel.LOW,
        rules=rules,
    )
    from dataclasses import replace

    return replace(base, quality=quality)


def _read_csv(path: Path) -> list[list[str]]:
    with path.open(encoding="utf-8-sig", newline="") as fh:
        return list(_csv.reader(fh))


def test_status_line_reports_excepted_count() -> None:
    status = quality_status(_doc().quality)
    assert status.excepted == 1
    assert "1 excepted" in quality_status_line(status)


def test_json_keeps_list_and_exception_fields(tmp_path: Path) -> None:
    (path,) = JsonWriter().write(_doc(), tmp_path / "out.json")
    payload = json.loads(path.read_text())
    issues = payload["quality"]["issues"]
    stale = next(i for i in issues if i["rule_id"] == "dq.stale_keyword.v1")
    assert stale["component_ids"] == ["evar1"]
    assert stale["occurrences"] == 1
    assert stale["excepted"] is True
    assert stale["exception_reason"] == "kept for audit trail"
    conv = next(i for i in issues if i["rule_id"] == "dq.naming_convention.v1")
    assert conv["excepted"] is False
    assert payload["quality"]["summary"]["excepted"] == 1


def test_csv_quality_issues_has_new_columns(tmp_path: Path) -> None:
    CsvWriter().write(_doc(), tmp_path / "out.csv")
    rows = _read_csv(tmp_path / "out.quality_issues.csv")
    header = rows[0]
    assert header[:7] == ["severity", "category", "type", "item_id", "item_name", "issue", "details"]
    assert header[7:] == ["rule_id", "component_ids", "occurrences", "exception"]
    by_rule = {r[7]: dict(zip(header, r, strict=True)) for r in rows[1:]}
    assert by_rule["dq.stale_keyword.v1"]["exception"] == "excepted (kept for audit trail)"
    assert by_rule["dq.stale_keyword.v1"]["component_ids"] == "evar1"
    assert by_rule["dq.naming_convention.v1"]["exception"] == ""
    assert by_rule["dq.naming_case_inconsistency.v1"]["component_ids"] == "evar2"


def test_excel_quality_sheet_has_new_columns(tmp_path: Path) -> None:
    target = tmp_path / "out.xlsx"
    ExcelWriter().write(_doc(), target)
    wb = load_workbook(target, read_only=True)
    try:
        rows = [tuple(r) for r in wb["Quality"].iter_rows(values_only=True)]
    finally:
        wb.close()
    header_idx = next(i for i, r in enumerate(rows) if r[0] == "severity")
    summary = {r[0]: r[1] for r in rows[:header_idx] if r[0]}
    assert summary["Excepted"] == 1
    header = rows[header_idx]
    assert header[7:11] == ("rule_id", "component_ids", "occurrences", "exception")
    stale = next(r for r in rows[header_idx + 1 :] if r[7] == "dq.stale_keyword.v1")
    assert stale[10] == "excepted (kept for audit trail)"


def test_markdown_lists_findings(tmp_path: Path) -> None:
    (path,) = MarkdownWriter().write(_doc(), tmp_path / "out.md")
    quality = path.read_text().split("## Quality\n", 1)[1].split("\n## ", 1)[0]
    assert "| Severity | Rule ID |" in quality
    assert "dq.naming_convention.v1" in quality
    assert "excepted (kept for audit trail)" in quality


def test_html_lists_findings(tmp_path: Path) -> None:
    (path,) = HtmlWriter().write(_doc(), tmp_path / "out.html")
    content = path.read_text()
    quality = content.split("<h2>Quality</h2>", 1)[1].split("</section>", 1)[0]
    assert "Rule ID" in quality
    assert "dq.stale_keyword.v1" in quality
    assert "excepted (kept for audit trail)" in quality


def test_quality_report_json_and_csv(tmp_path: Path) -> None:
    doc = _doc()
    issues = [_rehydrate_issue(d) for d in doc.quality["issues"]]
    summary = doc.quality["summary"]

    write_quality_report(issues=issues, summary=summary, target=tmp_path / "q.json", fmt="json")
    payload = json.loads((tmp_path / "q.json").read_text())
    assert payload["issues"] == doc.quality["issues"]
    assert payload["summary"]["excepted"] == 1

    write_quality_report(issues=issues, summary=summary, target=tmp_path / "q.csv", fmt="csv")
    rows = _read_csv(tmp_path / "q.csv")
    assert rows[0] == [
        "severity",
        "category",
        "type",
        "item_id",
        "item_name",
        "issue",
        "rule_id",
        "component_ids",
        "occurrences",
        "excepted",
        "exception_reason",
    ]
    stale = next(r for r in rows[1:] if r[6] == "dq.stale_keyword.v1")
    assert stale[7:] == ["evar1", "1", "true", "kept for audit trail"]


def test_legacy_issue_without_rule_id_renders_derived_rule(tmp_path: Path) -> None:
    from dataclasses import replace

    legacy = {
        "issues": [
            {
                "severity": "LOW",
                "category": "stale",
                "type": "version_suffix",
                "item_id": "evar3",
                "item_name": "x_v2",
                "issue": "m",
                "details": {},
            },
        ],
        "summary": {"by_severity": {"LOW": 1}, "total": 1, "verdict": "n/a"},
    }
    doc = replace(_doc(), quality=legacy)
    (path,) = MarkdownWriter().write(doc, tmp_path / "out.md")
    assert "dq.version_suffix.v1" in path.read_text()
    assert quality_status(legacy).excepted == 0
